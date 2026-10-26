# -*- coding: utf-8 -*-
"""LaTeX -> MathML (latex2mathml) -> OMML (Office MML2OMML.XSL) for native Word equations."""
import copy
from lxml import etree
import latex2mathml.converter

XSL = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
_xslt = etree.XSLT(etree.parse(XSL))


def latex_to_omml(tex):
    tex = tex.replace(r"\qquad", "\\text{\u2003\u2003}").replace(r"\quad", "\\text{\u2003}")
    mml = latex2mathml.converter.convert(tex)
    tree = etree.fromstring(mml.encode("utf-8"))
    om = _xslt(tree).getroot()
    return om                      # <m:oMath>


def add_equation(doc, tex, number=None, docx_para=None):
    """Append a centred display equation (optionally numbered '(n)' at the right margin via a tab)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm
    from docx.enum.text import WD_TAB_ALIGNMENT
    p = docx_para or doc.add_paragraph()
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    om = latex_to_omml(tex)
    p._p.append(copy.deepcopy(om))
    if number is not None:
        p.add_run(f"\t({number})")
    return p


def add_inline(p, tex):
    p._p.append(copy.deepcopy(latex_to_omml(tex)))


if __name__ == "__main__":
    import docx
    d = docx.Document()
    add_equation(d, r"h_i(t)=\ln H_{\mathrm{cs}}(t)+m(t)+\lambda_i a(t)+\sigma_i(t)+\varepsilon^{h}_i(t)", 1)
    add_equation(d, r"\mathbb{E}[h^{*}\mid h^{*}\le c]=\mu-\sqrt{S}\,\frac{\phi(\alpha)}{\Phi(\alpha)},\quad \alpha=\frac{c-\mu}{\sqrt{S}}", 2)
    add_equation(d, r"n_{d}=\left[\frac{12\,(z_{1-\alpha/2}+z_{1-\beta})^{2}\,s^{2}(1+\varphi)}{(1-\varphi)\,\delta^{2}}\right]^{1/3}", 3)
    d.save(r"C:\Users\User\AppData\Local\Temp\claude\d--Cluade-VsCode\dac3b60f-41c4-4c9e-8d8f-67c11baedcb3\scratchpad\eqtest.docx")
    print("ok")
