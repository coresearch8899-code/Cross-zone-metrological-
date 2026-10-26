# -*- coding: utf-8 -*-
"""Cover letter for Progress in Photovoltaics."""
import json, os
from docx import Document
from docx.shared import Pt, Cm
from ss_data import ROOT
import text_R14 as TX

R = json.load(open(os.path.join(ROOT, "results", "results_main.json")))
TH = json.load(open(os.path.join(ROOT, "results", "theory.json")))
s1 = R["sensors"]["zone1"]["rel3_drift_2025on_pct_yr"][0]
md = TH["mdd"]; fac = md["anchor_h1_clearsky"]["days_to_detect_89permil_yr"] / md["contrast_h1_h3"]["days_to_detect_89permil_yr"]

d = Document()
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11); st.paragraph_format.space_after = Pt(8)
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2.5)
P = lambda t, b=False: setattr(d.add_paragraph().add_run(t), "bold", b)
P("[Date]")
P("The Editor\nProgress in Photovoltaics: Research and Applications")
P("Dear Editor,")
P(f"Please consider the enclosed manuscript, \"{TX.TITLE}\", for publication as a Research Article in Progress in Photovoltaics.")
P("Every field indicator of PV performance (performance ratio, performance-loss rate, soiling loss) divides energy by a measured "
  "irradiance, and the field pyranometers that supply it drift and are rarely recalibrated. The manuscript asks what a PV system with "
  "several separately metered zones can establish about its own sensors and arrays without additional instruments, and how to estimate it. "
  "Its contributions are methodological:")
for b in ("A gauge (identifiability) theorem that classifies which monitoring indicators are identifiable from multi-zone data, with "
          "results on anchoring, on isolating a drifting sensor (three sensors are needed), on the minimum detectable drift and on the "
          "recalibration interval of a networked sensor.",
          "A metrological-soiling state-space model that estimates sensor sensitivities, array condition and intervention effects jointly "
          "and, through a Tobit Kalman update, uses PR values clipped by monitoring platforms as censored observations instead of discarding them.",
          "A semi-synthetic validation with the real sky and real truncation, showing that discarding clipped values biases drift estimates "
          "towards stability.",
          f"A demonstration on a 2.5 MWp three-zone system, where the method dated a reference drift of {s1:.0f} % per year, would have "
          f"raised an online alarm months before the end of the record, detected the drift about {fac:.0f} times faster than a clear-sky "
          "comparison, and supports staggered recalibration of networked sensors."):
    d.add_paragraph(b, style="List Bullet")
P("We are aware that the journal accepts systems papers only in exceptional cases. The manuscript is not a system report: the "
  "propositions hold for any number of zones, the estimator needs only data that commercial platforms already store, and the case "
  "study is used to demonstrate and test the method. We believe it is of interest to readers working on PV monitoring, performance-loss "
  "analysis and O&M.")
P("Disclosure of related submissions. An earlier manuscript on the same system was submitted to Solar Energy and declined because its "
  "scientific contribution was judged insufficient. The present manuscript is a new work: the theory, the estimator, the validation "
  "and the recalibration analysis are new, and the earlier event-study and cleaning-economics analyses have been removed. A separate "
  "manuscript using production data of the same installation, on biomass-burning haze and source attribution, is under consideration "
  "at another journal; it does not address sensor identifiability, censored estimation or recalibration. We would be glad to provide "
  "either manuscript to the editor on request.")
P("The work is original, has not been published and is not under consideration elsewhere. The author declares no conflicts of interest. "
  "Data and code are provided as supporting information.")
P("Yours sincerely,")
P("Thanapon Saengsuwan\nProgram in Energy Engineering and Electrical Technology, Faculty of Industrial Technology\n"
  "Chiang Rai Rajabhat University, Chiang Rai 57100, Thailand\nthanapon.sae@crru.ac.th")
fn = os.path.join(ROOT, "manuscript", "PIP", "Cover_Letter_PIP.docx")
d.save(fn); print("saved", fn)
